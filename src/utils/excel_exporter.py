"""
Excel 导出模块
负责将采集的数据导出为格式化的 Excel 文件
"""

from datetime import datetime
from pathlib import Path
from typing import List, Dict
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """Excel 导出器"""

    # 列定义（按 CLAUDE.md 规格）
    COLUMNS = {
        'keyword': 'keyword',
        'channel_title': 'KOL名称',
        'video_link': '视频链接',
        'view_count': '播放数量',
        'engagement_rate': '互动率',
        'contact_info': '联系方式',
        'like_count': '点赞数量',
        'comment_count': '评论数量',
        'video_title': '视频标题',
        'published_at': '发布日期',
        'duration': '视频时长',
        'subscriber_count': '频道订阅数',
        'channel_link': '频道链接',
        'description': '视频描述',
        'language': '视频语言'
    }

    def __init__(self, logger=None):
        """
        初始化 Excel 导出器

        Args:
            logger: 日志器（可选）
        """
        self.logger = logger

    def export(
        self,
        data: List[Dict],
        output_path: str,
        show_channel_link: bool = False,
        show_video_description: bool = False
    ) -> bool:
        """
        导出数据到 Excel

        Args:
            data: 数据列表
            output_path: 输出文件路径
            show_channel_link: 是否显示频道链接
            show_video_description: 是否显示视频描述

        Returns:
            bool: 是否导出成功
        """
        if not data:
            if self.logger:
                self.logger.warning("没有数据可导出")
            return False

        try:
            # 准备DataFrame
            df = self._prepare_dataframe(data, show_channel_link, show_video_description)

            # 导出到Excel
            df.to_excel(output_path, index=False, engine='openpyxl')

            # 格式化Excel
            self._format_excel(output_path, show_channel_link, show_video_description)

            if self.logger:
                self.logger.info(f"成功导出 {len(data)} 条数据到: {output_path}")

            return True

        except Exception as e:
            if self.logger:
                self.logger.error(f"导出 Excel 失败: {e}")
            return False

    def _prepare_dataframe(
        self,
        data: List[Dict],
        show_channel_link: bool,
        show_video_description: bool
    ) -> pd.DataFrame:
        """
        准备 DataFrame

        Args:
            data: 数据列表
            show_channel_link: 是否显示频道链接
            show_video_description: 是否显示视频描述

        Returns:
            pd.DataFrame: DataFrame
        """
        # 构建列顺序
        columns = [
            'keyword',
            'channel_title',
            'video_link',
            'view_count',
            'engagement_rate',
            'contact_info',
            'like_count',
            'comment_count',
            'video_title',
            'published_at',
            'duration',
            'subscriber_count'
        ]

        # 添加可选列
        if show_channel_link:
            columns.append('channel_link')
        if show_video_description:
            columns.append('description')

        # 添加视频语言
        columns.append('language')

        # 准备数据
        rows = []
        for item in data:
            row = {}
            for col in columns:
                row[col] = item.get(col, '')
            rows.append(row)

        # 创建 DataFrame，使用中文列名
        df = pd.DataFrame(rows, columns=columns)
        df.columns = [self.COLUMNS.get(col, col) for col in columns]

        return df

    def _format_excel(
        self,
        file_path: str,
        show_channel_link: bool,
        show_video_description: bool
    ):
        """
        格式化 Excel 文件

        Args:
            file_path: Excel 文件路径
            show_channel_link: 是否显示频道链接
            show_video_description: 是否显示视频描述
        """
        # 加载工作簿
        wb = load_workbook(file_path)
        ws = wb.active

        # 样式定义
        header_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')

        cell_font = Font(name='微软雅黑', size=10)
        cell_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        center_alignment = Alignment(horizontal='center', vertical='center')

        border = Border(
            left=Side(style='thin', color='D0D0D0'),
            right=Side(style='thin', color='D0D0D0'),
            top=Side(style='thin', color='D0D0D0'),
            bottom=Side(style='thin', color='D0D0D0')
        )

        # 设置表头样式
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        # 设置数据行样式
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.font = cell_font
                cell.border = border

                # 根据列类型设置对齐方式
                col_letter = get_column_letter(cell.column)
                if col_letter in ['D', 'E', 'G', 'H', 'L']:  # 数字列居中
                    cell.alignment = center_alignment
                else:
                    cell.alignment = cell_alignment

        # 设置列宽
        column_widths = {
            'A': 15,  # keyword
            'B': 20,  # KOL名称
            'C': 45,  # 视频链接
            'D': 12,  # 播放数量
            'E': 10,  # 互动率
            'F': 35,  # 联系方式
            'G': 12,  # 点赞数量
            'H': 12,  # 评论数量
            'I': 40,  # 视频标题
            'J': 12,  # 发布日期
            'K': 10,  # 视频时长
            'L': 12,  # 频道订阅数
        }

        col_index = ord('M')
        if show_channel_link:
            column_widths[chr(col_index)] = 45  # 频道链接
            col_index += 1
        if show_video_description:
            column_widths[chr(col_index)] = 50  # 视频描述
            col_index += 1

        column_widths[chr(col_index)] = 12  # 视频语言

        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width

        # 设置行高
        ws.row_dimensions[1].height = 25  # 表头行高
        for row in range(2, ws.max_row + 1):
            ws.row_dimensions[row].height = 20

        # 冻结首行
        ws.freeze_panes = 'A2'

        # 保存
        wb.save(file_path)

    @staticmethod
    def generate_filename(keyword: str, output_dir: str = '.') -> str:
        """
        生成导出文件名

        Args:
            keyword: 搜索关键词
            output_dir: 输出目录

        Returns:
            str: 完整文件路径
        """
        # 清理关键词（移除非法字符）
        safe_keyword = "".join(c for c in keyword if c.isalnum() or c in (' ', '-', '_')).strip()
        safe_keyword = safe_keyword[:50]  # 限制长度

        # 生成时间戳
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

        # 生成文件名
        filename = f"{safe_keyword}_{timestamp}.xlsx"

        # 返回完整路径
        return str(Path(output_dir) / filename)
